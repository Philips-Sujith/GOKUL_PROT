import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from pathlib import Path

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_ai_ml_guide():
    doc = Document()
    
    # Page Setup - Margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    # Styles
    # Primary Palette
    DARK_BLUE = RGBColor(15, 23, 42)    # #0f172a
    TEAL = RGBColor(13, 148, 136)       # #0d9488
    SLATE = RGBColor(71, 85, 105)       # #475569
    ROSE = RGBColor(225, 29, 72)        # #e11d48
    
    # Document Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_badge = title_p.add_run("USHNA KAAPPAAN — SCIENTIFIC HEAT EARLY WARNING SYSTEM\n")
    run_badge.font.name = "Arial"
    run_badge.font.size = Pt(10)
    run_badge.font.bold = True
    run_badge.font.color.rgb = TEAL

    run_title = title_p.add_run("Understanding AI & Machine Learning in Ushna Kaappaan")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = DARK_BLUE

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("A Complete, Step-by-Step Beginner’s Guide: From Core Concepts to Real-World Output")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = SLATE

    # Horizontal divider rule
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(14)
    run_div = p_div.add_run("―" * 60)
    run_div.font.color.rgb = RGBColor(203, 213, 225)

    # 1. Executive Summary & The Problem We Are Solving
    h1 = doc.add_heading("1. The Real-World Problem: Why Standard Temperature is Not Enough", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    
    doc.add_paragraph(
        "When you look at a weather app on your mobile phone, it tells you the ambient air temperature "
        "(for example, 34°C). However, any human being who has walked in direct sunlight at 2:00 PM on a humid day in "
        "Chennai or Madurai knows that 34°C feels radically different depending on whether the sun is beating down on your body, "
        "how humid the air is (which stops your sweat from evaporating), and whether there is any wind to cool you down."
    )
    
    doc.add_paragraph(
        "Furthermore, extreme heat does not merely cause discomfort; it is a major public health hazard that exacerbates "
        "cardiovascular disease, respiratory failure, kidney stress, and heatstroke, resulting in excess deaths, "
        "especially among outdoor laborers, the elderly, and vulnerable socioeconomic populations."
    )

    doc.add_paragraph(
        "Ushna Kaappaan solves this by combining two powerful paradigms:\n"
        "1. Physics-based Biometeorological Modeling (calculating how the human body exchanges heat with the atmosphere).\n"
        "2. Artificial Intelligence / Machine Learning (predicting state-level excess heat mortality risk based on "
        "environmental exposure and population vulnerability)."
    )

    # 2. Physics Engine vs. Machine Learning: The Two Pillars
    h1 = doc.add_heading("2. The Two Pillars: Physics Engine vs. Machine Learning", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "It is crucial for a beginner to understand the difference between the Physics Calculations and the Machine Learning Model "
        "used in this system:"
    )

    # Comparison Table
    table1 = doc.add_table(rows=1, cols=3)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table1.rows[0].cells
    hdr_cells[0].text = "Aspect"
    hdr_cells[1].text = "Pillar 1: Physics Engine (Thermofeel)"
    hdr_cells[2].text = "Pillar 2: Machine Learning Model"

    for cell in hdr_cells:
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 120, 120, 150, 150)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

    comp_data = [
        ("What is it?", "Exact mathematical & physical formulas (ECMWF thermofeel / Brode et al. 2012).", "Trained statistical AI algorithms (Scikit-Learn ensemble regressors)."),
        ("What does it calculate?", "Mean Radiant Temperature (MRT) & Universal Thermal Climate Index (UTCI) in °C.", "Daily Heat-Related Excess Mortality Rate per 100,000 population."),
        ("Spatial Resolution", "District-level: calculated individually for all 83 South Indian districts.", "State-level: computed aggregated across Tamil Nadu, Kerala, and Karnataka."),
        ("Why use this method?", "Human heat exchange obeys thermodynamics (radiation, evaporation, convection).", "Mortality is complex, non-linear, and depends on demographics, healthcare, and cumulative heat."),
        ("Nature of Output", "Continuous physical thermal stress index (°C) + 5 standard categories.", "Predicted deaths per 100k people + 4 Public Health Risk Bands (Baseline to Severe).")
    ]

    for row_idx, data in enumerate(comp_data):
        row_cells = table1.add_row().cells
        bg_color = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for i, text in enumerate(data):
            row_cells[i].text = text
            set_cell_background(row_cells[i], bg_color)
            set_cell_margins(row_cells[i], 100, 100, 140, 140)
            for p in row_cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.font.bold = True

    # 3. What is Machine Learning? (Beginner Concept)
    h1 = doc.add_heading("3. Machine Learning in Simple Terms", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "In traditional programming, humans write explicit rules: IF temperature > 38 THEN alert = True. "
        "However, real-world mortality from heat cannot be captured by a simple rule because it depends on multiple interacting variables: "
        "how long the heatwave has lasted, the percentage of elderly people, outdoor worker density, hospital access, and humidity."
    )
    doc.add_paragraph(
        "Machine Learning (ML) takes historical or calibrated data containing both the input features (weather, thermal stress, demographics) "
        "and the historical outcome (the mortality rate). The ML algorithm analyzes thousands of records, discovers the complex non-linear patterns, "
        "and produces a trained model. When live data is fed into this trained model, it accurately estimates the expected mortality risk."
    )

    # 4. The Step-by-Step ML Pipeline
    h1 = doc.add_heading("4. The Ushna Kaappaan ML Pipeline: Step-by-Step", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    # Step 4.1
    doc.add_heading("Step 4.1: Dataset & Calibration Baseline", level=2)
    doc.add_paragraph(
        "The mortality model is calibrated on synthetic epidemiological baseline data (Guin et al. 2025 statistical calibration framework) "
        "located at `dataset/mortality_synthetic_daily.csv`. It contains daily records spanning multiple years across South Indian states, "
        "integrating daily maximum temperatures, humidity, solar radiation, heatwave flags, and demographic profiles."
    )

    # Step 4.2
    doc.add_heading("Step 4.2: Feature Selection & Strict Zero Target Leakage", level=2)
    doc.add_paragraph(
        "A common mistake in beginner ML is 'Data Leakage' — giving the model information during training that it wouldn't have during live operation. "
        "For example, if you include raw death counts (like `heatstroke_deaths` or `heat_related_deaths_male`) in the input features, the model will simply "
        "add them up, creating a false 100% accuracy that completely fails in real life when future death counts are unknown!"
    )
    doc.add_paragraph(
        "Ushna Kaappaan enforces a strict Zero Target Leakage policy in `ml_pipeline/train_mortality_model.py`. All raw death counts, expected counts, "
        "and mortality anomalies are strictly excluded from the input predictors (X). Only valid environmental, thermal, and demographic features are used."
    )

    # Step 4.3
    doc.add_heading("Step 4.3: Chronological (Temporal) Train/Test Split", level=2)
    doc.add_paragraph(
        "In standard ML, data is often randomly shuffled. However, weather and health data are time series data. "
        "Randomly shuffling would allow the model to 'peek' into future days to guess the past. "
        "Ushna Kaappaan uses a 75% / 25% Chronological Split:\n"
        "• Train Set: 75% of the earliest historical dates.\n"
        "• Test Set: 25% of the most recent unseen future dates.\n"
        "This rigorously proves that the model can forecast future health risks on days it has never seen before."
    )

    # Step 4.4
    doc.add_heading("Step 4.4: Candidate Model Tournament", level=2)
    doc.add_paragraph(
        "Rather than guessing a single algorithm, our training script executes an automated tournament across 5 diverse algorithms:"
    )

    cand_table = doc.add_table(rows=1, cols=3)
    cand_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_hdr = cand_table.rows[0].cells
    c_hdr[0].text = "Algorithm"
    c_hdr[1].text = "Category / Mechanism"
    c_hdr[2].text = "Why It Was Evaluated"

    for cell in c_hdr:
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 120, 120, 150, 150)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

    candidates_info = [
        ("Ridge Regression", "Linear Model with L2 Regularization", "Fast linear baseline that prevents large unstable weights."),
        ("ElasticNet Regression", "Linear Model with L1 + L2 Regularization", "Combines feature selection (L1) with stability (L2)."),
        ("Random Forest Regressor", "Bagging Ensemble (150 Decision Trees)", "Averages multiple deep decision trees to model complex non-linear heat thresholds without overfitting."),
        ("Gradient Boosting Regressor", "Sequential Boosting Ensemble (120 Trees)", "Builds trees sequentially, correcting the residual errors of prior trees to capture subtle temperature-mortality curves."),
        ("Extra Trees Regressor", "Extremely Randomized Trees Ensemble", "Introduces random split points across trees to minimize variance on extreme weather spikes.")
    ]

    for idx, item in enumerate(candidates_info):
        r_cells = cand_table.add_row().cells
        bg = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        for i, text in enumerate(item):
            r_cells[i].text = text
            set_cell_background(r_cells[i], bg)
            set_cell_margins(r_cells[i], 100, 100, 140, 140)
            for p in r_cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.font.bold = True

    # 5. How Evaluation Metrics Work (Explained for Beginners)
    h1 = doc.add_heading("5. Understanding the Evaluation Metrics", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "When evaluating ML regression models, 4 key metrics are measured on the unseen Test set:\n"
        "1. RMSE (Root Mean Squared Error): The average error magnitude. Because errors are squared before averaging, large mistakes are penalized heavily. Lower is better.\n"
        "2. MAE (Mean Absolute Error): The direct average difference between predicted rate and actual rate. Lower is better.\n"
        "3. R² (Coefficient of Determination): Measures how much of the variance in mortality is successfully explained by the model (1.0 = perfect explanation, 0.0 = no better than guessing the average). Higher is better.\n"
        "4. Pearson Correlation: Measures how well the predicted trend tracks the actual spike in mortality (1.0 = perfect upward and downward synchronization)."
    )

    # 6. Live Inference Workflow: How an Output is Delivered
    h1 = doc.add_heading("6. Live Inference: How Output is Delivered in Production", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "When the Ushna Kaappaan application runs in real-time, the ML model functions through an automated 5-step pipeline:"
    )

    doc.add_paragraph(
        "Step 1: Telemetry Fetch — The backend fetches live weather from Open-Meteo for all 83 districts.\n"
        "Step 2: Physics Calculation — Thermofeel computes MRT and UTCI for each district.\n"
        "Step 3: State Feature Aggregation — District values are aggregated per state (Tamil Nadu, Kerala, Karnataka) and combined with demographic constants (e.g. elderly population %, outdoor worker fraction, healthcare access index).\n"
        "Step 4: Machine Learning Inference — The saved pipeline artifact (`mortality_model_pipeline.joblib`) performs instant mathematical inference, generating the predicted mortality rate (deaths per 100k residents per day).\n"
        "Step 5: Public Health Risk Band Mapping — The rate is mapped into actionable public health context:"
    )

    # Risk Bands Table
    band_table = doc.add_table(rows=1, cols=3)
    band_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    b_hdr = band_table.rows[0].cells
    b_hdr[0].text = "Risk Band"
    b_hdr[1].text = "Predicted Rate (Deaths / 100k / Day)"
    b_hdr[2].text = "Public Health & Administrative Directive"

    for cell in b_hdr:
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 120, 120, 150, 150)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

    band_data = [
        ("Baseline", "Rate ≤ 0.05", "Normal background levels. Standard routine health advisories apply."),
        ("Moderate Risk", "0.05 < Rate ≤ 0.15", "Elevated heat exposure stress. Activate cooling centers, distribute ORS at transit hubs."),
        ("High Risk", "0.15 < Rate ≤ 0.30", "Substantial excess heat mortality burden. Restrict peak-afternoon outdoor labor, alert PHC emergency wards."),
        ("Severe Risk", "Rate > 0.30", "Critical cumulative thermal emergency. Multi-agency escalation, hospital surge protocols, emergency water supply.")
    ]

    for idx, item in enumerate(band_data):
        r_cells = band_table.add_row().cells
        bg = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        for i, text in enumerate(item):
            r_cells[i].text = text
            set_cell_background(r_cells[i], bg)
            set_cell_margins(r_cells[i], 100, 100, 140, 140)
            for p in r_cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.font.bold = True

    # 7. Step-by-Step Example Walkthrough
    h1 = doc.add_heading("7. Real-World Walkthrough Example", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Let us follow a single data point from the sky to the final health alert:\n\n"
        "1. Open-Meteo Satellite & Weather Telemetry (Chennai at 1:00 PM):\n"
        "   • Air Temp: 35.0°C | Relative Humidity: 46% | Wind Speed: 2.5 m/s | Solar Irradiance: 930 W/m²\n\n"
        "2. Physics Engine (ECMWF Thermofeel):\n"
        "   • Radiation balance yields Mean Radiant Temperature (MRT) = 57.5°C (+22.5°C radiation load from direct sun).\n"
        "   • 6th-order multi-node heat balance yields UTCI = 40.5°C ('Very strong heat stress').\n\n"
        "3. State Aggregation & Demographic Synthesis (Tamil Nadu):\n"
        "   • Aggregated State UTCI: 39.8°C | Elderly population: 13.6% | Outdoor labor: 35% | Heat Action Plan index: 0.70.\n\n"
        "4. Machine Learning Model Inference:\n"
        "   • Gradient Boosting / Ensemble pipeline evaluates feature vector → Predicts 0.24 deaths per 100,000 population/day.\n\n"
        "5. Alert Engine & Multi-Channel Action:\n"
        "   • Rate falls in HIGH RISK band (0.15 - 0.30).\n"
        "   • GIS Map displays red choropleth on Chennai & neighboring districts.\n"
        "   • Admin Portal marks Chennai as a Priority Unit.\n"
        "   • Automated Telegram broadcast dispatches early-warning advisory to Municipalities, PHCs, and Worker Unions."
    )

    # Save document
    docs_dir = Path("docs")
    docs_dir.mkdir(parents=True, exist_ok=True)
    out_path = docs_dir / "Ushna_Kaappaan_AI_ML_Beginner_Guide.docx"
    doc.save(str(out_path))
    print(f"Successfully created: {out_path}")

if __name__ == "__main__":
    create_ai_ml_guide()
