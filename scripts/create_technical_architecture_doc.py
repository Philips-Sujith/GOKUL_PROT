import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from pathlib import Path

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_technical_architecture_doc():
    doc = Document()
    
    # Page Setup
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    DARK_BLUE = RGBColor(15, 23, 42)    # #0f172a
    TEAL = RGBColor(13, 148, 136)       # #0d9488
    SLATE = RGBColor(71, 85, 105)       # #475569
    
    # Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_badge = title_p.add_run("USHNA KAAPPAAN — SYSTEM ARCHITECTURE & ENGINEERING SPECIFICATION\n")
    run_badge.font.name = "Arial"
    run_badge.font.size = Pt(10)
    run_badge.font.bold = True
    run_badge.font.color.rgb = TEAL

    run_title = title_p.add_run("Technical Architecture, Algorithms, Models & Tech Stacks")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = DARK_BLUE

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Comprehensive Engineering Deep-Dive: Architecture Diagrams, Scientific Formulas, ML Pipelines & Workflow")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = SLATE

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(14)
    run_div = p_div.add_run("―" * 60)
    run_div.font.color.rgb = RGBColor(203, 213, 225)

    # 1. Complete Technology Stack
    h1 = doc.add_heading("1. Comprehensive Technology Stack", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Ushna Kaappaan is built as a robust, production-grade biometeorological monitoring and early warning system. "
        "The complete technology stack is segmented into distinct architectural layers:"
    )

    tech_table = doc.add_table(rows=1, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_hdr = tech_table.rows[0].cells
    t_hdr[0].text = "Layer"
    t_hdr[1].text = "Technology / Library"
    t_hdr[2].text = "Purpose & Role in Ushna Kaappaan"

    for cell in t_hdr:
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 120, 120, 150, 150)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

    tech_data = [
        ("Frontend UI Framework", "React 18 + TypeScript", "Type-safe, component-driven client application with responsive state management."),
        ("Styling & Icons", "TailwindCSS + Lucide React", "Design system with curated HSL palettes, responsive grid cards, and visual iconography."),
        ("GIS Mapping Engine", "Leaflet 1.9 + Leaflet-GeoJSON", "Client-side vector GIS engine rendering 83 contiguous district polygons with interactive selection."),
        ("Frontend Build Tool", "Vite 6.4 + PostCSS", "High-performance ESM bundler with hot module replacement and optimized production builds."),
        ("Backend Web API", "FastAPI (Python 3.14) + Uvicorn", "High-throughput asynchronous REST API with OpenAPI documentation and CORS management."),
        ("Data Modeling & ORM", "SQLAlchemy 2.0 + SQLite", "Declarative ORM schema modeling 83 districts, weather observations, alerts, and audit deliveries."),
        ("Data Validation", "Pydantic v2", "Strict runtime request/response payload validation, typing, and serialization."),
        ("Scientific Thermal Engine", "ECMWF Thermofeel + NumPy", "Vectorized biometeorological physics library calculating Solar Zenith, MRT, and UTCI."),
        ("Machine Learning", "Scikit-Learn (Joblib, Pandas)", "Automated regression pipeline (Random Forest / Gradient Boosting) predicting state mortality rates."),
        ("Weather Telemetry", "Open-Meteo REST API", "Direct radiation and multi-variable meteorological data ingestion without API key bottlenecks."),
        ("Broadcasting Channel", "Telegram Bot API", "Multi-role early warning dispatch engine with automatic MOCK/REAL transport switching."),
        ("Testing Suite", "Pytest + Pytest-Asyncio", "22 automated unit and integration tests covering alert engines, physics formulas, and ML inference.")
    ]

    for idx, item in enumerate(tech_data):
        r_cells = tech_table.add_row().cells
        bg = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        for i, text in enumerate(item):
            r_cells[i].text = text
            set_cell_background(r_cells[i], bg)
            set_cell_margins(r_cells[i], 100, 100, 140, 140)
            for p in r_cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.font.bold = True

    # 2. System Architecture & Component Interaction
    h1 = doc.add_heading("2. System Architecture & Data Flow", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "The system operates across three decoupled planes: Ingestion & Physics Plane, Storage & Intelligence Plane, "
        "and Presentation & Broadcast Plane."
    )

    arch_diagram_text = (
        "┌────────────────────────────────────────────────────────────────────────────────────────┐\n"
        "│                          1. INGESTION & SENSING PLANE                                  │\n"
        "│  [ Open-Meteo REST API ] ──(Every 2h Ingestion)──> 83 South India Districts Coordinates │\n"
        "│  (Temp, Humidity, Wind 10m, Surface Pressure, Direct & Diffuse Solar Radiation)        │\n"
        "└──────────────────────────────────────────┬─────────────────────────────────────────────┘\n"
        "                                           │\n"
        "                                           ▼\n"
        "┌────────────────────────────────────────────────────────────────────────────────────────┐\n"
        "│                          2. SCIENTIFIC & AI COMPUTATION ENGINE                         │\n"
        "│  [ ECMWF Thermofeel Physics Engine ]              [ Scikit-Learn ML Model Pipeline ]    │\n"
        "│  • Solar Zenith Angle (θz)                        • 75/25 Chronological Temporal Split │\n"
        "│  • Actual Vapour Pressure (Magnus)                • Zero Target Leakage Feature Matrix │\n"
        "│  • Mean Radiant Temperature (MRT) in °C           • State Demographic Scaling          │\n"
        "│  • UTCI 6th-Order Polynomial Multi-Node (°C)      • Predicted Mortality Rate / 100,000 │\n"
        "└──────────────────────────────────────────┬─────────────────────────────────────────────┘\n"
        "                                           │\n"
        "                                           ▼\n"
        "┌────────────────────────────────────────────────────────────────────────────────────────┐\n"
        "│                          3. DATABASE & PERSISTENCE LAYER                               │\n"
        "│  [ SQLite via SQLAlchemy ORM ]                                                         │\n"
        "│  • District Records (83 units) • Observations Archive • Active Alerts • Audit Trail    │\n"
        "└──────────────────────────────────────────┬─────────────────────────────────────────────┘\n"
        "                                           │\n"
        "                                           ▼\n"
        "┌────────────────────────────────────────────────────────────────────────────────────────┐\n"
        "│                          4. REST API & DISPATCH GATEWAYS                               │\n"
        "│  [ FastAPI 0.110+ Asynchronous Router Engine ]                                         │\n"
        "│  • /api/districts • /api/alerts • /api/mortality/overview • /api/admin/escalate        │\n"
        "└─────────────────────┬──────────────────────────────────────────┬───────────────────────┘\n"
        "                      │                                          │\n"
        "                      ▼                                          ▼\n"
        "┌──────────────────────────────────────────────┐ ┌───────────────────────────────────────┐\n"
        "│          CLIENT PRESENTATION LAYER           │ │        BROADCAST ALERT GATEWAY        │\n"
        "│  • React 18 + Leaflet Vector Choropleth      │ │  • Telegram Bot API (REAL Mode)       │\n"
        "│  • Live 83-District Spatial Telemetry        │ │  • Transparent Mock Audit Transport   │\n"
        "│  • 4 Population Profiles & Health Advisories │ │  • Multi-Role Target Dispatch (PHCs,  │\n"
        "│  • Admin Console & Priority Triage           │ │    Municipalities, Worker Unions)     │\n"
        "└──────────────────────────────────────────────┘ └───────────────────────────────────────┘"
    )

    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(6)
    p_box.paragraph_format.space_after = Pt(10)
    run_box = p_box.add_run(arch_diagram_text)
    run_box.font.name = "Consolas"
    run_box.font.size = Pt(7.5)
    run_box.font.color.rgb = DARK_BLUE

    # 3. Scientific Mathematical Formulas & Algorithms
    h1 = doc.add_heading("3. Core Mathematical Formulas & Physical Algorithms", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    # 3.1 Solar Zenith Angle
    doc.add_heading("3.1 Solar Position & Zenith Angle (θz)", level=2)
    doc.add_paragraph(
        "To calculate direct solar radiation absorption on the human body, the position of the sun in the sky "
        "is computed for each district's exact latitude (φ), longitude, and day of year (N):\n"
        "• Solar Declination (δ): δ = 23.45° · sin( (360/365) · (284 + N) )\n"
        "• Hour Angle (ω): ω = 15° · (Solar_Time - 12)\n"
        "• Cosine of Zenith Angle: cos(θz) = sin(φ)·sin(δ) + cos(φ)·cos(δ)·cos(ω)\n"
        "When cos(θz) ≤ 0, the sun is below the horizon (nighttime), and solar radiation is clamped to zero."
    )

    # 3.2 Actual Vapour Pressure
    doc.add_heading("3.2 Actual Vapour Pressure (ea) via Magnus Equation", level=2)
    doc.add_paragraph(
        "Human evaporative heat loss through sweating is governed by water vapour pressure in the surrounding air. "
        "From air temperature (Ta in °C) and Relative Humidity (RH in %):\n"
        "• Saturation Vapour Pressure: e_sat(Ta) = 6.112 · exp( (17.67 · Ta) / (Ta + 243.5) )  [hPa]\n"
        "• Actual Vapour Pressure: e_a = (RH / 100) · e_sat(Ta)  [hPa]\n"
        "This metric is directly consumed by the multi-node thermal equations."
    )

    # 3.3 Mean Radiant Temperature (MRT)
    doc.add_heading("3.3 Mean Radiant Temperature (MRT)", level=2)
    doc.add_paragraph(
        "MRT represents the uniform temperature of an imaginary enclosure in which the radiant heat transfer from the human body "
        "equals the radiant heat transfer in the actual non-uniform environment (including direct sunlight, sky diffuse radiation, "
        "and ground thermal reflection).\n\n"
        "Formula (Plane Irradiance Radiation Balance):\n"
        "T_mrt = [ ( (1/σ) · ( Σ(w_i · E_i) + (1 - α_p)/ε_p · f_p · I_dir + (1 - α_p)/ε_p · Σ(f_i · D_i) ) ) ]^(0.25) - 273.15\n\n"
        "Where:\n"
        "• σ = Stefan-Boltzmann constant (5.670374 × 10⁻⁸ W/m²·K⁴)\n"
        "• I_dir = Direct normal beam solar irradiance\n"
        "• D_i = Diffuse sky & ground reflected radiation\n"
        "• f_p = Projected area factor of a standing human body as a function of solar zenith angle θz\n"
        "• α_p = Human skin & clothing albedo (~0.30)\n"
        "• ε_p = Human skin emissivity (~0.97)\n"
        "In direct South Indian sunshine, MRT regularly exceeds ambient air temperature by +20°C to +25°C."
    )

    # 3.4 UTCI Polynomial
    doc.add_heading("3.4 Universal Thermal Climate Index (UTCI)", level=2)
    doc.add_paragraph(
        "UTCI is defined as the equivalent ambient temperature (°C) of a reference environment that would elicit the same physiological "
        "response (sweat rate, core body temperature, skin wettedness, blood flow) in the Fiala multi-node dynamic human thermal model.\n\n"
        "The ECMWF thermofeel engine implements the authoritative 6th-order multi-variate polynomial regression (Brode et al. 2012):\n"
        "UTCI(Ta, Tmrt, v_10m, ea) = Ta + Offset(ΔTmrt, v_10m, ea, Ta)\n\n"
        "Where Offset is a polynomial of degree 6 containing over 200 interaction terms capturing non-linear convective cooling, "
        "radiant warming, and humidity-induced sweat inhibition."
    )

    # 3.5 IDW Interpolation
    doc.add_heading("3.5 Spatial Interpolation: Inverse Distance Weighting (IDW)", level=2)
    doc.add_paragraph(
        "For the optional continuous spatial thermal field visualization, the system calculates estimated UTCI at any grid coordinate (x, y) "
        "using Inverse Distance Weighting from all 83 district observation stations (x_i, y_i):\n"
        "UTCI_interp(x, y) = [ Σ (w_i · UTCI_i) ] / [ Σ w_i ]\n"
        "Where w_i = 1 / ( (dist( (x,y), (x_i, y_i) ) + ε)^p ), with power p = 2.0."
    )

    # 4. Machine Learning Model Architecture & Training
    h1 = doc.add_heading("4. Machine Learning Architecture, Training & Hyperparameters", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "The Machine Learning component is implemented in `ml_pipeline/train_mortality_model.py`. "
        "Key architectural parameters and configurations include:"
    )

    ml_spec_table = doc.add_table(rows=1, cols=2)
    ml_spec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    m_hdr = ml_spec_table.rows[0].cells
    m_hdr[0].text = "Hyperparameter / Specification"
    m_hdr[1].text = "Engineering Value & Rationale"

    for cell in m_hdr:
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 120, 120, 150, 150)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

    ml_specs = [
        ("Target Variable (y)", "heat_related_mortality_rate_per_100000 (Daily excess deaths per 100k population)"),
        ("Model Pipeline Imputer", "SimpleImputer(strategy='median') for missing value robustness"),
        ("Feature Scaling", "StandardScaler() applied to linear candidates (Ridge, ElasticNet)"),
        ("Winning Regressor", "GradientBoostingRegressor / RandomForestRegressor (Ensemble of 120–150 trees)"),
        ("Tree Hyperparameters", "max_depth=4–8, min_samples_split=4, learning_rate=0.05, random_state=42"),
        ("Train/Test Split Strategy", "75% Train / 25% Test strictly divided chronologically by calendar date"),
        ("Target Leakage Prevention", "21 target-correlated columns (death totals, components, anomalies) hard-excluded"),
        ("Serialization Format", "Joblib binary pipeline artifact: `mortality_model_pipeline.joblib`"),
        ("Inference Speed", "< 2.5 milliseconds per state inference vector")
    ]

    for idx, item in enumerate(ml_specs):
        r_cells = ml_spec_table.add_row().cells
        bg = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        for i, text in enumerate(item):
            r_cells[i].text = text
            set_cell_background(r_cells[i], bg)
            set_cell_margins(r_cells[i], 100, 100, 140, 140)
            for p in r_cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.font.bold = True

    # 5. End-to-End Operational Workflow
    h1 = doc.add_heading("5. End-to-End Operational Workflow", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "The system executes an automated, continuous 6-stage lifecycle:\n\n"
        "1. Ingestion Cycle (Every 2 Hours):\n"
        "   The background scheduler triggers `run_full_pipeline()`. It connects to Open-Meteo REST API, "
        "   fetching 2m air temperature, 10m wind speed, relative humidity, surface pressure, and downwelling solar radiation "
        "   for all 83 district geographic coordinates.\n\n"
        "2. Physics & Thermal Computation:\n"
        "   For every district, Thermofeel calculates Solar Zenith, Vapour Pressure, MRT, and UTCI. "
        "   The district is assigned its standard category (No Stress, Moderate, Strong, Very Strong, Extreme).\n\n"
        "3. Database Archival & Time-Series History:\n"
        "   Observations are committed to SQLite via SQLAlchemy. Time-series historical records are appended to enable "
        "   trend analysis and multi-day heat accumulation monitoring.\n\n"
        "4. Machine Learning State Risk Assessment:\n"
        "   State-level feature aggregations are combined with demographic constants and passed to the ML pipeline, "
        "   evaluating the excess daily mortality rate per 100k population for Tamil Nadu, Kerala, and Karnataka.\n\n"
        "5. Automated Alert Evaluation & Threshold Triage:\n"
        "   The alert engine evaluates operational thresholds:\n"
        "   • UTCI ≥ 46°C → Operational Severity: EXTREME\n"
        "   • UTCI ≥ 38°C → Operational Severity: SEVERE\n"
        "   • UTCI ≥ 32°C → Operational Severity: HIGH\n\n"
        "6. Multi-Channel Presentation & Broadcast Delivery:\n"
        "   • Public Portal: Real-time 83-district interactive choropleth map updates with custom demographic advisories.\n"
        "   • Admin Portal: High-stress districts are populated into the triage console with one-click escalation controls.\n"
        "   • Telegram Broadcaster: Structured early-warning notifications are dispatched to municipal and health channels."
    )

    # Save document
    docs_dir = Path("docs")
    docs_dir.mkdir(parents=True, exist_ok=True)
    out_path = docs_dir / "Ushna_Kaappaan_Technical_Architecture_and_Engineering_Deep_Dive.docx"
    doc.save(str(out_path))
    print(f"Successfully created: {out_path}")

if __name__ == "__main__":
    create_technical_architecture_doc()
